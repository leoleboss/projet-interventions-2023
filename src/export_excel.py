import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties


def exporter_excel(df, top_departements, categories):
    chemin_sortie = "outputs/reporting_interventions_2023.xlsx"

    colonne_region = [col for col in df.columns if "gion" in col][0]
    colonne_departement = [col for col in df.columns if "partement" in col][0]

    regions_uniques = ["Toutes"] + sorted(df[colonne_region].dropna().unique())
    categories_uniques = [
        "Toutes",
        "Incendies",
        "Secours à personne",
        "Accidents",
        "Opérations diverses"
    ]

    departements_uniques = sorted(df[colonne_departement].dropna().unique())

    filtres_df = pd.DataFrame({
        "regions": pd.Series(regions_uniques),
        "categories": pd.Series(categories_uniques)
    })

    calcul_df = pd.DataFrame({
        "departement": departements_uniques
    })

    with pd.ExcelWriter(chemin_sortie, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="donnees_nettoyees", index=False)
        filtres_df.to_excel(writer, sheet_name="listes_filtres", index=False)
        calcul_df.to_excel(writer, sheet_name="calcul_dashboard", index=False)

    workbook = load_workbook(chemin_sortie)

    dashboard = workbook.create_sheet("tableau_de_bord", 0)
    feuille_filtres = workbook["listes_filtres"]
    feuille_calcul = workbook["calcul_dashboard"]
    feuille_donnees = workbook["donnees_nettoyees"]

    vert = "6AA84F"
    vert_clair = "D9EAD3"
    gris_clair = "F3F3F3"

    titre_font = Font(size=18, bold=True, color="FFFFFF")
    sous_titre_font = Font(size=11, bold=True, color="FFFFFF")
    kpi_font = Font(size=12, bold=True)
    valeur_kpi_font = Font(size=15, bold=True)

    remplissage_vert = PatternFill("solid", fgColor=vert)
    remplissage_vert_clair = PatternFill("solid", fgColor=vert_clair)
    remplissage_gris = PatternFill("solid", fgColor=gris_clair)

    bordure = Border(
        left=Side(style="thin", color=vert),
        right=Side(style="thin", color=vert),
        top=Side(style="thin", color=vert),
        bottom=Side(style="thin", color=vert)
    )

    for col in range(1, 15):
        dashboard.column_dimensions[chr(64 + col)].width = 16

    dashboard.merge_cells("A1:N2")
    dashboard["A1"] = "Tableau de bord - Interventions des secours en 2023"
    dashboard["A1"].fill = remplissage_vert
    dashboard["A1"].font = titre_font
    dashboard["A1"].alignment = Alignment(horizontal="center", vertical="center")

    dashboard.merge_cells("A4:N5")
    dashboard["A4"] = "Zone de filtres"
    dashboard["A4"].fill = remplissage_vert_clair
    dashboard["A4"].alignment = Alignment(horizontal="center", vertical="center")

    dashboard["B6"] = "Filtre région"
    dashboard["C6"] = "Toutes"

    dashboard["E6"] = "Filtre catégorie"
    dashboard["F6"] = "Toutes"

    for cellule in ["B6", "E6"]:
        dashboard[cellule].fill = remplissage_vert
        dashboard[cellule].font = Font(bold=True, color="FFFFFF")
        dashboard[cellule].alignment = Alignment(horizontal="center")
        dashboard[cellule].border = bordure

    for cellule in ["C6", "F6"]:
        dashboard[cellule].fill = remplissage_gris
        dashboard[cellule].alignment = Alignment(horizontal="center")
        dashboard[cellule].border = bordure

    validation_region = DataValidation(
        type="list",
        formula1=f"'listes_filtres'!$A$2:$A${feuille_filtres.max_row}",
        allow_blank=False
    )

    validation_categorie = DataValidation(
        type="list",
        formula1="'listes_filtres'!$B$2:$B$6",
        allow_blank=False
    )

    dashboard.add_data_validation(validation_region)
    dashboard.add_data_validation(validation_categorie)

    validation_region.add(dashboard["C6"])
    validation_categorie.add(dashboard["F6"])

    critere_region = 'IF($C$6="Toutes","*",$C$6)'

    # Colonnes Excel utiles
    # C = région
    # E = département
    # R = incendies
    # AM = secours à personne
    # AS = accidents de circulation
    # BR = opérations diverses
    # BS = total interventions

    feuille_calcul["B1"] = "incendies"
    feuille_calcul["C1"] = "secours_personne"
    feuille_calcul["D1"] = "accidents"
    feuille_calcul["E1"] = "operations_diverses"
    feuille_calcul["F1"] = "total_filtre"

    for ligne in range(2, feuille_calcul.max_row + 1):
        departement_cellule = f"A{ligne}"

        feuille_calcul[f"B{ligne}"] = (
            f'=SUMIFS(donnees_nettoyees!$R$2:$R$200,'
            f'donnees_nettoyees!$E$2:$E$200,{departement_cellule},'
            f'donnees_nettoyees!$C$2:$C$200,IF(tableau_de_bord!$C$6="Toutes","*",tableau_de_bord!$C$6))'
        )

        feuille_calcul[f"C{ligne}"] = (
            f'=SUMIFS(donnees_nettoyees!$AM$2:$AM$200,'
            f'donnees_nettoyees!$E$2:$E$200,{departement_cellule},'
            f'donnees_nettoyees!$C$2:$C$200,IF(tableau_de_bord!$C$6="Toutes","*",tableau_de_bord!$C$6))'
        )

        feuille_calcul[f"D{ligne}"] = (
            f'=SUMIFS(donnees_nettoyees!$AS$2:$AS$200,'
            f'donnees_nettoyees!$E$2:$E$200,{departement_cellule},'
            f'donnees_nettoyees!$C$2:$C$200,IF(tableau_de_bord!$C$6="Toutes","*",tableau_de_bord!$C$6))'
        )

        feuille_calcul[f"E{ligne}"] = (
            f'=SUMIFS(donnees_nettoyees!$BR$2:$BR$200,'
            f'donnees_nettoyees!$E$2:$E$200,{departement_cellule},'
            f'donnees_nettoyees!$C$2:$C$200,IF(tableau_de_bord!$C$6="Toutes","*",tableau_de_bord!$C$6))'
        )

        feuille_calcul[f"F{ligne}"] = (
            f'=IF(tableau_de_bord!$F$6="Toutes",'
            f'SUM(B{ligne}:E{ligne}),'
            f'IF(tableau_de_bord!$F$6="Incendies",B{ligne},'
            f'IF(tableau_de_bord!$F$6="Secours à personne",C{ligne},'
            f'IF(tableau_de_bord!$F$6="Accidents",D{ligne},'
            f'IF(tableau_de_bord!$F$6="Opérations diverses",E{ligne},0)))))'
        )

    dashboard["A9"] = "Top 10 départements"
    dashboard["A9"].fill = remplissage_vert
    dashboard["A9"].font = sous_titre_font

    dashboard["A10"] = "Département"
    dashboard["B10"] = "Interventions"

    for i in range(1, 11):
        ligne_dashboard = 10 + i
        dashboard[f"B{ligne_dashboard}"] = f'=LARGE(calcul_dashboard!$F$2:$F$200,{i})'
        dashboard[f"A{ligne_dashboard}"] = (
            f'=INDEX(calcul_dashboard!$A$2:$A$200,'
            f'MATCH(B{ligne_dashboard},calcul_dashboard!$F$2:$F$200,0))'
        )

    for cellule in ["A10", "B10"]:
        dashboard[cellule].fill = remplissage_gris
        dashboard[cellule].font = Font(bold=True)
        dashboard[cellule].border = bordure

    dashboard["D9"] = "Grandes catégories"
    dashboard["D9"].fill = remplissage_vert
    dashboard["D9"].font = sous_titre_font

    dashboard["D10"] = "Catégorie"
    dashboard["E10"] = "Interventions"

    categories_lignes = [
        ("Incendies", "R"),
        ("Secours à personne", "AM"),
        ("Accidents", "AS"),
        ("Opérations diverses", "BR")
    ]

    for index, (nom, colonne_excel) in enumerate(categories_lignes, start=11):
        dashboard.cell(row=index, column=4).value = nom
        dashboard.cell(row=index, column=5).value = (
            f'=IF($F$6="Toutes",'
            f'SUMIFS(donnees_nettoyees!${colonne_excel}$2:${colonne_excel}$200,'
            f'donnees_nettoyees!$C$2:$C$200,{critere_region}),'
            f'IF($F$6=D{index},'
            f'SUMIFS(donnees_nettoyees!${colonne_excel}$2:${colonne_excel}$200,'
            f'donnees_nettoyees!$C$2:$C$200,{critere_region}),0))'
        )

    for cellule in ["D10", "E10"]:
        dashboard[cellule].fill = remplissage_gris
        dashboard[cellule].font = Font(bold=True)
        dashboard[cellule].border = bordure

    kpis = {
        "J9": ("Total interventions", '=SUM(E11:E14)'),
        "L9": ("Département n°1", '=A11'),
        "J13": ("Part secours à personne", '=IF(J10=0,0,E12/J10)'),
        "L13": ("Moyenne / département", '=IF(COUNTIF(B11:B20,">0")=0,0,J10/COUNTIF(B11:B20,">0"))')
    }

    for position, (nom, valeur) in kpis.items():
        dashboard[position] = nom
        dashboard[position].fill = remplissage_vert_clair
        dashboard[position].font = kpi_font
        dashboard[position].alignment = Alignment(horizontal="center")
        dashboard[position].border = bordure

        cellule_valeur = dashboard.cell(
            row=dashboard[position].row + 1,
            column=dashboard[position].column
        )
        cellule_valeur.value = valeur
        cellule_valeur.font = valeur_kpi_font
        cellule_valeur.alignment = Alignment(horizontal="center")
        cellule_valeur.border = bordure

    dashboard["J14"].number_format = "0.0%"

    graphique_top = BarChart()
    graphique_top.type = "bar"
    graphique_top.style = 10
    graphique_top.title = "Top 10 départements"
    graphique_top.y_axis.title = ""
    graphique_top.x_axis.title = ""
    graphique_top.legend = None
    graphique_top.height = 8
    graphique_top.width = 18

    data_top = Reference(dashboard, min_col=2, min_row=10, max_row=20)
    labels_top = Reference(dashboard, min_col=1, min_row=11, max_row=20)

    graphique_top.add_data(data_top, titles_from_data=True)
    graphique_top.set_categories(labels_top)

    graphique_top.dLbls = DataLabelList()
    graphique_top.dLbls.showVal = True
    graphique_top.dLbls.showCatName = True
    graphique_top.dLbls.showSerName = False
    graphique_top.dLbls.showLegendKey = False
    graphique_top.dLbls.numFmt = '#,##0'

    dashboard.add_chart(graphique_top, "A25")

    graphique_categories = PieChart()
    graphique_categories.title = "Répartition des interventions"
    graphique_categories.height = 8
    graphique_categories.width = 14

    data_categories = Reference(dashboard, min_col=5, min_row=10, max_row=14)
    labels_categories = Reference(dashboard, min_col=4, min_row=11, max_row=14)

    graphique_categories.add_data(data_categories, titles_from_data=True)
    graphique_categories.set_categories(labels_categories)

    graphique_categories.dLbls = DataLabelList()
    graphique_categories.dLbls.showPercent = True
    graphique_categories.dLbls.showVal = False
    graphique_categories.dLbls.showCatName = False
    graphique_categories.dLbls.showSerName = False
    graphique_categories.dLbls.showLegendKey = False
    graphique_categories.legend.position = "b"

    graphique_categories.dLbls.txPr = RichText(
        p=[
            Paragraph(
                pPr=ParagraphProperties(
                    defRPr=CharacterProperties(
                        solidFill="FFFFFF",
                        b=True
                    )
                )
            )
        ]
    )

    dashboard.add_chart(graphique_categories, "I25")

    format_nombre = '#,##0'

    for cellule in ["J10", "L14"]:
        dashboard[cellule].number_format = format_nombre

    for ligne in range(11, 21):
        dashboard[f"B{ligne}"].number_format = format_nombre

    for ligne in range(11, 15):
        dashboard[f"E{ligne}"].number_format = format_nombre

    for ligne in feuille_donnees.iter_rows(min_row=2):
        for cellule in ligne:
            if isinstance(cellule.value, (int, float)):
                cellule.number_format = format_nombre

    feuille_donnees.auto_filter.ref = feuille_donnees.dimensions

    feuille_filtres.sheet_state = "hidden"
    feuille_calcul.sheet_state = "hidden"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    workbook.save(chemin_sortie)

    print("Fichier Excel créé avec tableau de bord dynamique :", chemin_sortie)
    